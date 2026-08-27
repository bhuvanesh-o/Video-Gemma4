# -*- coding: utf-8 -*-
"""
step_4_prompt_document.py

Runs BOTH models on the single best truck crop step_3 saved for each logged
truck:
    1. Gemma 4 26B A4B via OpenRouter (cloud)
    2. Gemma 4 E2B via LiteRT-LM (local, no internet needed for inference)

Writes 16 new columns total: raw output + 7 fields, per model, clearly
suffixed so they never collide.
"""
import os
import time
import json
import pandas as pd
from openpyxl import load_workbook

from step_4_functions_prompt_document import (
    get_client, call_gemma_api, call_litert_e2b, extract_json, extract_fields, FIELD_LABELS,
)
from storage import dir_for
from progress import emit

CORE_PROMPT = (
    "You are analyzing one image of a truck captured by a traffic/CCTV camera. "
    "Look closely and answer the following parameters exactly, in this structured JSON format:\n\n"
    "1. VEHICLE TYPE: Specify the exact type of truck (e.g., container truck, tanker, flatbed, "
    "mini-truck, pickup, trailer, dump truck, etc.). If not a truck, state what it actually is.\n"
    "2. COLOR: State the primary body color, and secondary color if visible (e.g., cabin vs container).\n"
    "3. NUMBER PLATE: Read and transcribe the license/number plate text exactly as visible. "
    "If partially visible or unreadable, state 'partially visible: <what you can read>' or 'not readable'. "
    "Do not guess characters you cannot see clearly.\n"
    "4. BRAND / MAKE: If any manufacturer logo, brand name, or markings are visible (e.g., Tata, Ashok Leyland, "
    "Volvo, Eicher), state it. Otherwise say 'not visible'.\n"
    "5. LOAD / CARGO: Describe what the truck appears to be carrying, if anything is visible.\n"
    "6. CONDITION/DAMAGE: Note any visible damage, dents, rust, or unusual markings.\n"
    "7. ADDITIONAL DETAILS: Any other identifying features — stickers, text on the body, company name, etc.\n\n"
    "Be precise and objective. If a detail is not visible, explicitly say so instead of guessing."

    """
    IMPORTANT OUTPUT FORMAT:

    Return ONLY one valid JSON object.

    Use EXACTLY these keys:

    {
    "VEHICLE TYPE": "",
    "COLOR": "",
    "NUMBER PLATE": "",
    "BRAND / MAKE": "",
    "LOAD / CARGO": "",
    "CONDITION/DAMAGE": "",
    "ADDITIONAL DETAILS": ""
    }

    Do not write any explanation before the JSON.
    Do not write any explanation after the JSON.
    Do not use bullet points.
    Do not return Markdown.
    Do not wrap the JSON in ```json code fences.
    Always return all seven keys.
    If information is unavailable, use "not visible".
    """
    
)

# ADDED: this strict JSON-schema enforcement block (exact key names,
# "no markdown/no code fences/no explanation") was appended after real
# E2B LiteRT runs showed less reliable JSON adherence than 26B A4B's
# response_format={"type": "json_object"} parameter provides natively —
# LiteRT-LM's CLI has no equivalent structured-output guarantee, so the
# prompt itself has to do that enforcement instead.



RAW_26B_COLUMN = "26B A4B Raw JSON"
RAW_LITERT_COLUMN = "E2B LiteRT Raw Output"
FIELD_LABELS_26B = [f"{label} (26B A4B)" for label in FIELD_LABELS]
FIELD_LABELS_LITERT = [f"{label} (E2B LiteRT)" for label in FIELD_LABELS]


def run_26b_a4b(client, image_path):
    parsed = call_gemma_api(client, image_path, CORE_PROMPT)
    fields = extract_fields(parsed)
    result = {f"{k} (26B A4B)": v for k, v in fields.items()}
    result[RAW_26B_COLUMN] = json.dumps(parsed, ensure_ascii=False, indent=2)
    return result



def run_e2b_litert(image_path, retries=3):

    """
    Wraps call_litert_e2b() (ONE subprocess call) with ITS OWN retry loop —
    two separate layers of retry logic, worth distinguishing: this function
    retries by calling call_litert_e2b() fresh, in full, up to `retries`
    times, specifically when JSON PARSING fails (not when the subprocess
    itself errors — that still propagates up immediately, uncaught here).
    26B A4B's equivalent retry logic instead lives INSIDE call_gemma_api()
    itself — the two models' retry strategies aren't structured identically,
    which is fine, just worth knowing when debugging either path.
    """

    last_raw_text = ""

    for attempt in range(1, retries + 1):
        print(f"[E2B] Attempt {attempt}/{retries}")
        raw_text = call_litert_e2b(
            image_path,
            CORE_PROMPT
        )

        last_raw_text = raw_text

        try:
            parsed = extract_json(raw_text)
            fields = extract_fields(parsed)
            result = {
                f"{k} (E2B LiteRT)": v
                for k, v in fields.items()
            }
            result[RAW_LITERT_COLUMN] = raw_text
            return result
        except (ValueError, json.JSONDecodeError) as e:

            print(f"[E2B] JSON parse failed: {e}")
            print("[E2B] Raw response:")
            print(raw_text)

            if attempt < retries:
                time.sleep(1)


    # All attempts failed
    result = {
        f"{k} (E2B LiteRT)": "N/A"
        for k in FIELD_LABELS
    }

    result[RAW_LITERT_COLUMN] = (
        f"[JSON parsing failed after {retries} attempts]\n\n"
        f"Last raw output:\n{last_raw_text}"
    )

    return result



def blank_result(raw_col, field_cols, message):
    result = {label: "N/A" for label in field_cols}
    result[raw_col] = message
    return result


def main(job_id):
    ASSET_DIR = dir_for(job_id, "final_assets")
    TRUCK_CROP_DIR = os.path.join(ASSET_DIR, "truck_crops")
    EXCEL_PATH = os.path.join(ASSET_DIR, "Vehicle_Registry_Master.xlsx")

    if not os.path.exists(EXCEL_PATH):
        emit("No registry found — skipping AI descriptions.", job_id=job_id, stage="describing", ui_message="")
        return

    df = pd.read_excel(EXCEL_PATH)
    if df.empty:
        emit("Registry is empty — skipping AI descriptions.", job_id=job_id, stage="describing", ui_message="")
        return

    api_key = os.environ.get("OPENROUTER_API_KEY")
    client = get_client(api_key) if api_key else None
    if not api_key:
        emit("⚠️ OPENROUTER_API_KEY not set — 26B A4B analysis will be skipped.", job_id=job_id,
             stage="describing", ui_message="")

    total = len(df)
    all_row_results = []

    # ==========================================================================
    # STEP 4 BENCHMARK STATISTICS
    # ==========================================================================

    step4_start = time.perf_counter()

    stats = {
        "26b_total_sec": 0.0,
        "26b_success": 0,
        "26b_failed": 0,

        "e2b_total_sec": 0.0,
        "e2b_success": 0,
        "e2b_failed": 0,

        "images_analyzed": 0,
        "images_missing": 0,
    }

    for idx, row in df.iterrows():
        vehicle_id = str(row["Vehicle ID"])
        truck_img_path = os.path.join(TRUCK_CROP_DIR, f"{vehicle_id}_truck.jpg")
        row_result = {}

        if not os.path.exists(truck_img_path):

            stats["images_missing"] += 1


            row_result.update(blank_result(RAW_26B_COLUMN, FIELD_LABELS_26B, "No image available"))
            row_result.update(blank_result(RAW_LITERT_COLUMN, FIELD_LABELS_LITERT, "No image available"))
            all_row_results.append(row_result)
            continue

        emit(f"🖼️ Analyzing {vehicle_id} ({idx + 1}/{total})...", job_id=job_id,
             stage="describing", ui_message=f"Analyzing {vehicle_id} ({idx + 1}/{total})...",
             meta={"clip_progress": {"current": idx + 1, "total": total}})

        # ----------------------------------------------------------------------
        # Pass 1 — 26B A4B OpenRouter
        # ----------------------------------------------------------------------

        if client:

            start_26b = time.perf_counter()

            try:
                result_26b = run_26b_a4b(
                    client,
                    truck_img_path
                )

                row_result.update(result_26b)
                stats["26b_success"] += 1

            except Exception as e:
                stats["26b_failed"] += 1

                row_result.update(
                    blank_result(
                        RAW_26B_COLUMN,
                        FIELD_LABELS_26B,
                        f"[analysis failed: {e}]"
                    )
                )

                emit(
                    f"⚠️ 26B A4B failed for {vehicle_id}: {e}",
                    job_id=job_id,
                    stage="describing",
                    ui_message=""
                )

            finally:
                stats["26b_total_sec"] += (
                    time.perf_counter() - start_26b
                )

        else:
            row_result.update(
                blank_result(
                    RAW_26B_COLUMN,
                    FIELD_LABELS_26B,
                    "API key not set"
                )
            )

        # ----------------------------------------------------------------------
        # Pass 2 — E2B LiteRT
        # ----------------------------------------------------------------------

        start_e2b = time.perf_counter()

        try:
            result_e2b = run_e2b_litert(
                truck_img_path
            )

            row_result.update(result_e2b)

            # run_e2b_litert may return normally even if all JSON retries failed,
            # so inspect its raw-output column before counting it as successful.
            raw_e2b = result_e2b.get(
                RAW_LITERT_COLUMN,
                ""
            )

            if raw_e2b.startswith("[JSON parsing failed"):
                stats["e2b_failed"] += 1
            else:
                stats["e2b_success"] += 1

        except Exception as e:
            stats["e2b_failed"] += 1

            row_result.update(
                blank_result(
                    RAW_LITERT_COLUMN,
                    FIELD_LABELS_LITERT,
                    f"[analysis failed: {e}]"
                )
            )

            emit(
                f"⚠️ E2B LiteRT failed for {vehicle_id}: {e}",
                job_id=job_id,
                stage="describing",
                ui_message=""
            )

        finally:
            stats["e2b_total_sec"] += (
                time.perf_counter() - start_e2b
            )

        all_row_results.append(row_result)
        time.sleep(1)  # gentle on OpenRouter's free-tier rate limit

    # --------------------------------------------------------------------------
    # WRITE ALL NEW COLUMNS via openpyxl — preserves existing hyperlink formulas
    # --------------------------------------------------------------------------
    all_new_columns = [RAW_26B_COLUMN] + FIELD_LABELS_26B + [RAW_LITERT_COLUMN] + FIELD_LABELS_LITERT

    wb = load_workbook(EXCEL_PATH, data_only=False)
    ws = wb.active
    headers = {cell.value: cell.column for cell in ws[1]}

    col_for = {}
    for col_name in all_new_columns:
        if col_name in headers:
            col_for[col_name] = headers[col_name]
        else:
            new_col_idx = ws.max_column + 1
            ws.cell(row=1, column=new_col_idx, value=col_name)
            headers[col_name] = new_col_idx
            col_for[col_name] = new_col_idx

    for row_idx, result in enumerate(all_row_results, start=2):
        for col_name in all_new_columns:
            ws.cell(row=row_idx, column=col_for[col_name], value=result.get(col_name, "N/A"))

    wb.save(EXCEL_PATH)
    emit(f"✅ AI descriptions added for {total} vehicles (both models).", job_id=job_id,
         stage="describing", ui_message="Descriptions added to your registry.")


    # ==========================================================================
    # FINAL STEP 4 BENCHMARK RESULTS
    # ==========================================================================

    stats["step4_total_sec"] = (
        time.perf_counter() - step4_start
    )

    # Number of actual 26B requests attempted
    attempted_26b = (
        stats["26b_success"]
        + stats["26b_failed"]
    )

    # Number of actual E2B requests attempted
    attempted_e2b = (
        stats["e2b_success"]
        + stats["e2b_failed"]
    )

    # Average latency per image
    stats["26b_avg_sec_per_image"] = (
        stats["26b_total_sec"] / attempted_26b
        if attempted_26b > 0
        else 0
    )

    stats["e2b_avg_sec_per_image"] = (
        stats["e2b_total_sec"] / attempted_e2b
        if attempted_e2b > 0
        else 0
    )

    # Success rates
    stats["26b_success_rate_pct"] = (
        100 * stats["26b_success"] / attempted_26b
        if attempted_26b > 0
        else 0
    )

    stats["e2b_success_rate_pct"] = (
        100 * stats["e2b_success"] / attempted_e2b
        if attempted_e2b > 0
        else 0
    )

    print("\n========== STEP 4 BENCHMARK ==========")

    print(
        f"26B A4B   : "
        f"{stats['26b_total_sec']:.2f}s total | "
        f"{stats['26b_avg_sec_per_image']:.2f}s/image | "
        f"{stats['26b_success']}/{attempted_26b} success | "
        f"{stats['26b_success_rate_pct']:.1f}%"
    )

    print(
        f"E2B LiteRT: "
        f"{stats['e2b_total_sec']:.2f}s total | "
        f"{stats['e2b_avg_sec_per_image']:.2f}s/image | "
        f"{stats['e2b_success']}/{attempted_e2b} success | "
        f"{stats['e2b_success_rate_pct']:.1f}%"
    )

    print(f"Images analyzed : {stats['images_analyzed']}")
    print(f"Images missing  : {stats['images_missing']}")
    print(f"Step 4 total    : {stats['step4_total_sec']:.2f}s")
    print("======================================\n")

    return stats

if __name__ == "__main__":
    main(job_id="local_test")